#!/usr/bin/env python3
"""Render an interview source log from interviews.json into Word, and emit
source-log rows that /productdd can cite.

    python3 build_interview_log.py interviews.json --out COMPANY_interviews.docx
    python3 build_interview_log.py interviews.json --out COMPANY_interviews.docx --sources sources.json

interviews.json:
{
  "company": "Acme",
  "interviews": [
    {
      "id": "I-01",
      "name": "Jane Doe",                     // or "Anonymised — Head of Logistics"
      "company": "Retailer X",
      "role": "Head of Supply Chain",
      "date": "2026-09-02",
      "relationship": "current customer | churned | prospect | ex-employee | expert | partner",
      "channel": "call | in person | written",
      "edited_transcript": "Cleaned Q&A text. Speaker labels Q:/A:.",
      "key_quotes": ["Short verbatim lines worth citing."],
      "takeaways": ["What it changes in the memo."],
      "modules": ["PMF evidence", "Monetization"],
      "confidence": "[B]"
    }
  ]
}

Why a script for this: transcripts are long and the useful part is the
reference block (who, when, in what capacity) — the thing that lets a
partner check a claim. The renderer standardises that block so every
interview reads the same and the source log builds itself.
"""

import argparse
import json
import os

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
ACCENT = RGBColor(0x0F, 0x3D, 0x5C)
FONT = "Calibri"


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def run(p, text, size=8.5, bold=False, italic=False, color=INK):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = FONT
    return r


def para(doc, text, **kw):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(kw.pop("space_after", 3))
    run(p, text, **kw)
    return p


def build(data, out):
    doc = Document()
    s = doc.sections[0]
    s.top_margin = s.bottom_margin = Inches(0.6)
    s.left_margin = s.right_margin = Inches(0.7)
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(8.5)

    para(doc, f"{data.get('company', '')} \u2014 Interview log", size=16, bold=True, color=ACCENT)
    para(doc, f"{len(data.get('interviews', []))} interviews \u00b7 edited transcripts \u00b7 "
              "quotes marked verbatim", size=8, italic=True, color=MUTED, space_after=8)

    # index table
    t = doc.add_table(rows=1, cols=6)
    t.style = "Table Grid"
    for i, h in enumerate(["ID", "Name", "Company", "Role", "Date", "Relationship"]):
        c = t.rows[0].cells[i]
        shade(c, "DCE5EC")
        c.paragraphs[0].text = ""
        run(c.paragraphs[0], h, size=8, bold=True)
    for iv in data.get("interviews", []):
        cells = t.add_row().cells
        for i, k in enumerate(["id", "name", "company", "role", "date", "relationship"]):
            cells[i].paragraphs[0].text = ""
            run(cells[i].paragraphs[0], str(iv.get(k, "")), size=8, bold=(i == 0))

    planned = data.get("planned", [])
    if planned:
        para(doc, "Planned — not yet conducted", size=9, bold=True, color=ACCENT, space_after=2)
        for pl in planned:
            para(doc, f"\u25aa {pl.get('name', '')} \u2014 {pl.get('role', '')}, {pl.get('company', '')}. "
                      f"Why: {pl.get('why', '')}", size=8, color=MUTED)

    docs_src = data.get("documents", [])
    if docs_src:
        para(doc, "Documents received", size=9, bold=True, color=ACCENT, space_after=2)
        for d in docs_src:
            para(doc, f"\u25aa {d.get('id', '')} \u2014 {d.get('what', '')} ({d.get('date', '')}) "
                      f"{d.get('finding', '')}", size=8, color=MUTED)

    for iv in data.get("interviews", []):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        para(doc, f"{iv.get('id', '')} \u2014 {iv.get('name', '')}", size=12, bold=True, color=ACCENT,
             space_after=1)
        para(doc, f"{iv.get('company', '')} \u00b7 {iv.get('role', '')} \u00b7 {iv.get('date', '')} \u00b7 "
                  f"{iv.get('relationship', '')} \u00b7 {iv.get('channel', '')} \u00b7 "
                  f"confidence {iv.get('confidence', '[B]')}",
             size=8, italic=True, color=MUTED, space_after=6)

        if iv.get("takeaways"):
            para(doc, "What it changes", size=9, bold=True, color=ACCENT, space_after=2)
            for x in iv["takeaways"]:
                para(doc, "\u25aa " + x)
        if iv.get("key_quotes"):
            para(doc, "Verbatim", size=9, bold=True, color=ACCENT, space_after=2)
            for q in iv["key_quotes"]:
                para(doc, "\u201c" + q + "\u201d", italic=True)
        if iv.get("modules"):
            para(doc, "Feeds: " + ", ".join(iv["modules"]), size=8, color=MUTED, space_after=6)
        para(doc, "Edited transcript", size=9, bold=True, color=ACCENT, space_after=2)
        for line in (iv.get("edited_transcript") or "").split("\n"):
            if not line.strip():
                continue
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            if line.startswith(("Q:", "Q :")):
                run(p, line, bold=True)
            else:
                run(p, line)

    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    doc.save(out)
    return out


def source_rows(data):
    """Rows for the productdd source-log exhibit: claim placeholder, source, date, tag."""
    rows = []
    for d in data.get("documents", []):
        rows.append([f"Document {d.get('id', '')}", d.get("what", ""), d.get("date", ""),
                     d.get("confidence", "[A]")])
    for iv in data.get("interviews", []):
        rows.append([
            f"Interview {iv.get('id', '')}",
            f"{iv.get('name', '')}, {iv.get('role', '')}, {iv.get('company', '')} "
            f"({iv.get('relationship', '')})",
            iv.get("date", ""),
            iv.get("confidence", "[B]"),
        ])
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("interviews_json")
    ap.add_argument("--out", required=True)
    ap.add_argument("--sources", help="also write source-log rows as JSON here")
    a = ap.parse_args()
    with open(a.interviews_json) as f:
        data = json.load(f)
    print(f"Wrote {build(data, a.out)}")
    if a.sources:
        with open(a.sources, "w") as f:
            json.dump({"columns": ["Claim", "Source", "Date", "Tag"],
                       "rows": source_rows(data)}, f, indent=2)
        print(f"Wrote {a.sources}")


if __name__ == "__main__":
    main()
