#!/usr/bin/env python3
"""onbrand path A — re-render a source doc.json through a target's brand theme.

    python3 apply_theme.py --doc-json ACME_productdd_doc.json \
        --builder .../productdd/scripts/build_product_doc.py \
        --theme themes/example-brand.json \
        --prepared-for "Meridian Partners" --by "Your Name" \
        --out ACME_productdd_MERIDIAN.docx

Patches the renderer's palette constants in a temporary copy and runs it, so charts, bands,
score ramps and shading are all generated in-brand rather than recoloured afterwards. The
original skill is never touched.

    --fonts safe|brand   safe (default) uses the theme's installed fallback faces
    --anonymise          replace the subject company with "Company <initial>"
    --alias "A=B"        additional replacements, repeatable
"""

import argparse

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from make_theme import dist, hexs, lstar, mix_on_white  # noqa: E402

WHITE = "#FFFFFF"


# ------------------------------------------------------------------ palette patch
def rgbcolor(h):
    h = hexs(h)
    return f"RGBColor(0x{h[0:2]}, 0x{h[2:4]}, 0x{h[4:6]})"


def build_constants(t, fonts_mode):
    c, s, f = t["colors"], t["semantic"], t["font"]
    body = f["brand"] if fonts_mode == "brand" else f["fallback"]
    # bands use the separated variant; severity TEXT keeps the dark one (it sits on white)
    sb = t.get("semantic_band") or s
    verdict = {}
    for k in ("invest", "back", "proceed"):
        verdict[k] = hexs(sb["success"])
    for k in ("conditional", "conditional invest", "hold", "watch"):
        verdict[k] = hexs(sb["warning"])
    for k in ("pass", "decline", "no"):
        verdict[k] = hexs(sb["error"])

    return {
        # name            -> (kind, value)
        "INK":         ("rgb", c["ink"]),
        "MUTED":       ("rgb", c["muted"]),
        "WHITE":       ("rgb", WHITE),
        "ACCENT":      ("rgb", c["accent"]),
        "ACCENT_HEX":  ("hex", c["accent"]),
        "BAND_HEX":    ("hex", c["band_fill"]),
        "CALLOUT_HEX": ("hex", c["callout"]),
        "HEADER_HEX":  ("hex", c["header_fill"]),
        "ZEBRA_HEX":   ("hex", c["zebra"]),
        "POS_HEX":     ("hex", mix_on_white(s["success"], 0.16)),
        "NEG_HEX":     ("hex", mix_on_white(s["error"], 0.16)),
        "NEU_HEX":     ("hex", mix_on_white(c["muted"], 0.12)),
        "BODY_FONT":   ("str", body),
        "CHART_COLORS": ("list", t["chart"]),
        "VERDICT_HEX": ("dict", verdict),
        "SCORE_HEX":   ("intdict", {int(k): hexs(v) for k, v in t["score_ramp"].items()}),
        "SEVERITY_HEX": ("dict", {"high": hexs(s["error"]), "medium": hexs(s["warning"]),
                                  "med": hexs(s["warning"]), "low": hexs(c["muted"])}),
    }


def render_value(kind, val):
    if kind == "rgb":
        return rgbcolor(val)
    if kind == "hex":
        return f'"{hexs(val)}"'
    if kind == "str":
        return f'"{val}"'
    if kind == "list":
        return "[" + ", ".join(f'"#{hexs(v)}"' for v in val) + "]"
    if kind == "dict":
        return "{" + ", ".join(f'"{k}": "{v}"' for k, v in val.items()) + "}"
    if kind == "intdict":
        return "{" + ", ".join(f'{k}: "{v}"' for k, v in val.items()) + "}"
    raise ValueError(kind)


ASSIGN = re.compile(r"^(?P<name>[A-Z_]+)\s*=\s*(?P<body>.*?)(?=^\s*(?:[A-Z_]+\s*=|def |class |#|$))",
                    re.M | re.S)


def patch_palette(src, theme, fonts_mode, log):
    consts = build_constants(theme, fonts_mode)
    out, seen = src, set()

    def repl(m):
        name = m.group("name")
        if name not in consts or name in seen:
            return m.group(0)
        seen.add(name)
        kind, val = consts[name]
        log.append(f"  {name} -> {render_value(kind, val)[:70]}")
        return f"{name} = {render_value(kind, val)}\n"

    out = ASSIGN.sub(repl, out, count=0)
    missing = [k for k in consts if k not in seen]
    if missing:
        log.append(f"  (renderer has no {', '.join(missing)} — skipped)")

    # hairline default argument
    out = re.sub(r'(def hairlines\([^)]*color=")[0-9A-Fa-f]{6}(")',
                 lambda m: m.group(1) + hexs(theme["colors"]["hairline"]) + m.group(2), out)
    return out, seen


# ------------------------------------------------- stray literals outside the palette
LITERAL_RGB = re.compile(r"RGBColor\(0x([0-9A-Fa-f]{2}),\s*0x([0-9A-Fa-f]{2}),\s*0x([0-9A-Fa-f]{2})\)")


def palette_for_matching(theme):
    c = theme["colors"]
    pool = [c["ink"], c["muted"], c["accent"], c["accent_fill"], c["band_fill"],
            c["band_text"], c["header_fill"], c["zebra"], c["callout"], WHITE]
    # on-band subtitle tints: the renderers use pale blues for text sitting on the header band
    pool += [mix_on_white(c["band_text"], 1.0), mix_on_white(c["accent_fill"], 0.55)]
    seen, out = set(), []
    for p in pool:
        if p and p.upper() not in seen:
            seen.add(p.upper()); out.append(p.upper())
    return out


def nearest_preserving_lightness(h, pool):
    """Nearest theme colour, weighting lightness heavily so text stays legible where it was."""
    return min(pool, key=lambda p: abs(lstar(p) - lstar(h)) * 12 + dist(p, h) * 0.05)


def remap_stray(src, theme, log, skip_before):
    pool = palette_for_matching(theme)
    head, tail = src[:skip_before], src[skip_before:]

    def repl(m):
        h = "#" + "".join(m.group(i).upper() for i in (1, 2, 3))
        if h in ("#FFFFFF", "#000000"):
            return m.group(0)
        n = nearest_preserving_lightness(h, pool)
        log.append(f"  stray {h} -> {n}")
        return rgbcolor(n)

    return head + LITERAL_RGB.sub(repl, tail)


# ------------------------------------------------------------------ doc.json edits
def walk_strings(obj, fn):
    if isinstance(obj, str):
        return fn(obj)
    if isinstance(obj, list):
        return [walk_strings(v, fn) for v in obj]
    if isinstance(obj, dict):
        return {k: walk_strings(v, fn) for k, v in obj.items()}
    return obj


def apply_aliases(doc, aliases):
    if not aliases:
        return doc
    pairs = sorted(aliases.items(), key=lambda kv: -len(kv[0]))

    def sub(s):
        for a, b in pairs:
            s = re.sub(re.escape(a), b, s, flags=re.I)
        return s

    return walk_strings(doc, sub)


# ------------------------------------------------------------------ brand pass
# Colour alone does not make a document look like a brand. A brand is recognised by its
# display face, its case, its tracking and where it puts its signature colour. The palette
# patch above handles hues; this pass handles identity. It runs on the rendered .docx, so it
# works for any renderer whose output has a shaded header band and ruled section headings.

def _fill_of(cell):
    from docx.oxml.ns import qn
    tcPr = cell._tc.tcPr
    if tcPr is None:
        return None
    shd = tcPr.find(qn("w:shd"))
    return (shd.get(qn("w:fill")) or "").upper() if shd is not None else None


def _set_tracking(run, em, size_pt):
    """Word stores character spacing in twentieths of a point, as an absolute value —
    so a brand's -2% em tracking has to be resolved against the run's own size."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    twips = int(round(em * size_pt * 20))
    if twips == 0:
        return
    rPr = run._r.get_or_add_rPr()
    for old in rPr.findall(qn("w:spacing")):
        rPr.remove(old)
    el = OxmlElement("w:spacing")
    el.set(qn("w:val"), str(twips))
    rPr.append(el)


def _caps(run):
    """Display in capitals via w:caps, not by rewriting the text — copy-paste out of the
    document then still yields the original casing."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rPr = run._r.get_or_add_rPr()
    if rPr.find(qn("w:caps")) is None:
        rPr.append(OxmlElement("w:caps"))


def _style(run, *, font=None, color=None, size=None, bold=None, upper=False, tracking=None):
    from docx.shared import Pt, RGBColor
    if upper and run.text.strip():
        _caps(run)
    if font:
        run.font.name = font
    if color:
        run.font.color.rgb = RGBColor.from_string(hexs(color))
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if tracking:
        pt = (run.font.size.pt if run.font.size else 9)
        _set_tracking(run, tracking, pt)


def _border(el, edge, color, sz=24, space=0):
    """sz is in eighths of a point: 24 = 3pt."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pPr = el.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.append(bdr)
    for old in bdr.findall(qn(f"w:{edge}")):
        bdr.remove(old)
    b = OxmlElement(f"w:{edge}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), hexs(color))
    bdr.append(b)


def _cell_top_border(cell, color, sz=24):
    """tcPr children are schema-ordered: tcBorders must sit BEFORE w:shd or the border is
    silently dropped — which is how a rule can vanish from an otherwise correct document."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    bdrs = tcPr.find(qn("w:tcBorders"))
    if bdrs is None:
        bdrs = OxmlElement("w:tcBorders")
        shd = tcPr.find(qn("w:shd"))
        if shd is not None:
            shd.addprevious(bdrs)
        else:
            tcPr.append(bdrs)
    for old in bdrs.findall(qn("w:top")):
        bdrs.remove(old)
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(sz))
    top.set(qn("w:color"), hexs(color))
    bdrs.append(top)


def _shade(cell, color):
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is not None:
        shd.set(qn("w:fill"), hexs(color))


def brand_pass(doc, theme, ty, display, body, log):
    """Four moves, in the order a reader notices them."""
    c = theme["colors"]
    sem = theme["semantic"]
    upper = ty.get("display_case") == "upper"
    tr = ty.get("display_tracking") or 0
    verdict_fills = {hexs(v) for v in
                     list(sem.values()) + list((theme.get("semantic_band") or {}).values())}

    band_done = False
    for tbl in doc.tables:
        if len(tbl.rows) != 1 or len(tbl.columns) != 1:
            continue
        cell = tbl.rows[0].cells[0]
        fill = _fill_of(cell)

        # 1. The header band — the company name becomes a wordmark, the classification
        #    takes the signature colour as TEXT (how most brands actually use a loud hue),
        #    and the meta line drops back so the two above it can be read at a glance.
        if not band_done and fill == hexs(c["band_fill"]):
            paras = cell.paragraphs
            runs = [r for r in (paras[0].runs if paras else []) if r.text.strip()]
            if runs:
                _style(runs[0], font=display, upper=upper, tracking=tr,
                       color=c["band_text"], bold=True)
                log.append("  header band: title set in " + display +
                           (" (uppercase)" if upper else ""))
            for r in runs[1:]:
                _style(r, font=body, color=c["on_band_accent"], bold=True)
            for p in paras[1:]:
                for r in p.runs:
                    _style(r, font=body, color=c["on_band_muted"])
            band_done = True
            continue

        # 2. The verdict band — keep the brand's own system colour and let a rule in the
        #    signature colour do the separating. Shifting the hue to force a luminance gap
        #    invents a colour the company does not own.
        if fill in verdict_fills:
            call = None
            for role, v in sem.items():
                if hexs(v) in verdict_fills and hexs((theme.get("semantic_band") or {})
                                                     .get(role, v)) == fill:
                    call = role
            if call:
                _shade(cell, sem[call])
                log.append(f"  verdict band: restored true {call} {sem[call]} + "
                           f"{c['rule']} rule")
            _cell_top_border(cell, c["rule"], sz=24)
            runs = [r for r in (cell.paragraphs[0].runs if cell.paragraphs else [])
                    if r.text.strip()]
            if runs:
                _style(runs[0], font=display, upper=upper, tracking=tr)

    # 3. Section headings — the renderer rules them in the accent colour. The signature
    #    colour carries no text there, so it can take the line, which is what puts the
    #    brand on every page instead of only page one.
    from docx.oxml.ns import qn
    from docx.shared import Pt
    ruled = 0
    for p in doc.paragraphs:
        pPr = p._p.pPr
        if pPr is None:
            continue
        bdr = pPr.find(qn("w:pBdr"))
        if bdr is None:
            continue
        bottom = bdr.find(qn("w:bottom"))
        if bottom is None or (bottom.get(qn("w:color")) or "").upper() != hexs(c["accent"]):
            continue
        bottom.set(qn("w:color"), hexs(c["rule"]))
        bottom.set(qn("w:sz"), "12")
        # A little more air above each rule, but only a little: every point added here is
        # multiplied by the number of headings and can push a whole extra page.
        cur = p.paragraph_format.space_before
        p.paragraph_format.space_before = Pt(max(9, cur.pt if cur else 0))
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            _style(r, font=display, upper=upper, tracking=tr, color=c["accent"])
        ruled += 1
    if ruled:
        log.append(f"  {ruled} section headings: {display} + {c['rule']} rule, more air above")

    # 4. Standalone display lines (exhibit dividers and the like)
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.size and r.font.size.pt >= 13 and r.font.bold:
                _style(r, font=display, upper=upper, tracking=tr)


def decorate(path, theme, prepared_for, by, date_line, fonts_mode, log):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    c = theme["colors"]
    ty = theme.get("type") or {}
    if fonts_mode == "brand":
        body = ty.get("body") or theme["font"]["brand"]
        display = ty.get("display") or body
    else:
        body = ty.get("body_fallback") or theme["font"]["fallback"]
        display = ty.get("display_fallback") or body

    doc = Document(path)
    brand_pass(doc, theme, ty, display, body, log)

    # attribution: small, tracked out, sitting on the signature rule
    bits = [f"Prepared for {prepared_for}" if prepared_for else None,
            by, date_line, "Unsolicited work sample"]
    line = "   ·   ".join(b for b in bits if b).upper()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(line)
    _style(r, font=body, color=c["accent"], size=6.5, bold=True, tracking=0.08)
    _border(p._p, "bottom", c["rule"], sz=12, space=3)
    doc.element.body.insert(0, p._p)

    footer = doc.sections[0].footer
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fr = fp.add_run(" ".join(x for x in [by, "·", f"prepared for {prepared_for}"] if x))
    _style(fr, font=body, color=c["muted"], size=6.5)
    sp = fp.add_run("\t\t")
    sp.font.size = Pt(6.5)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)

    doc.core_properties.author = by or ""
    doc.core_properties.title = os.path.basename(path)
    doc.core_properties.comments = (f"Prepared for {prepared_for}. Palette and type derived "
                                    f"from {theme.get('source','their site')}; no trademarked "
                                    f"marks used.")
    doc.save(path)


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--doc-json", required=True)
    ap.add_argument("--builder", required=True, help="the source skill's build_*.py")
    ap.add_argument("--theme", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--fonts", choices=["safe", "brand"], default="safe")
    ap.add_argument("--prepared-for", default="")
    ap.add_argument("--by", default="")
    ap.add_argument("--date", default="")
    ap.add_argument("--anonymise", action="store_true")
    ap.add_argument("--alias", action="append", default=[], metavar="FROM=TO")
    ap.add_argument("--keep-temp", action="store_true")
    a = ap.parse_args()

    theme = json.load(open(a.theme))
    doc = json.load(open(a.doc_json))

    aliases = {}
    if a.anonymise:
        company = doc.get("company", "")
        if company:
            aliases[company] = f"Company {company[0].upper()}"
    for pair in a.alias:
        k, _, v = pair.partition("=")
        aliases[k.strip()] = v.strip()
    doc = apply_aliases(doc, aliases)

    src = open(a.builder).read()
    log = []
    patched, seen = patch_palette(src, theme, a.fonts, log)
    if not seen:
        sys.exit(f"No palette constants found in {a.builder}. Either it is not one of the "
                 f"memo renderers, or its constant names changed — use retheme_docx.py instead.")
    # remap literals below the palette block only
    anchor = max((patched.find(f"\n{n} =") for n in seen), default=0)
    end = patched.find("\n", anchor + 1)
    patched = remap_stray(patched, theme, log, end if end > 0 else 0)

    tmp = tempfile.mkdtemp(prefix="onbrand_")
    bdir = os.path.dirname(os.path.abspath(a.builder))
    bpath = os.path.join(tmp, "themed_builder.py")
    open(bpath, "w").write(patched)
    djson = os.path.join(tmp, "doc.json")
    json.dump(doc, open(djson, "w"))

    os.makedirs(os.path.dirname(os.path.abspath(a.out)) or ".", exist_ok=True)
    env = dict(os.environ, PYTHONPATH=bdir + os.pathsep + env_path())
    res = subprocess.run([sys.executable, bpath, djson, "--out", a.out],
                         capture_output=True, text=True, env=env)
    if res.returncode != 0:
        print(res.stdout); print(res.stderr, file=sys.stderr)
        sys.exit(f"renderer failed — patched copy kept at {bpath}")

    decorate(a.out, theme, a.prepared_for, a.by, a.date, a.fonts, log)

    print(f"theme: {theme.get('brand')}  ({a.fonts} fonts)")
    for l in log:
        print(l)
    if aliases:
        print("  aliases: " + ", ".join(f"{k} -> {v}" for k, v in aliases.items()))
    print(f"wrote {a.out}")
    if not a.keep_temp:
        shutil.rmtree(tmp, ignore_errors=True)
    else:
        print(f"patched renderer: {bpath}")


def env_path():
    return os.environ.get("PYTHONPATH", "")


if __name__ == "__main__":
    main()
