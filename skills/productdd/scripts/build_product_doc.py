#!/usr/bin/env python3
"""Render a product due-diligence memo (or a value-creation plan) into styled Word.

Usage:
    python3 build_product_doc.py doc.json --out /path/to/COMPANY_productdd.docx

Every section is optional: omit a key and it is skipped. The same renderer serves
/productdd (verdict + scorecard + evidence) and /valuecreation (levers + 90-day plan),
because the two documents differ in content, not in typography.

See references/product-schema.md for the full structure.
"""

import argparse
import json
import os
import tempfile

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ---------------------------------------------------------------- palette
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x0F, 0x3D, 0x5C)
ACCENT_HEX = "0F3D5C"
CALLOUT_HEX = "EEF3F7"
HEADER_HEX = "DCE5EC"
ZEBRA_HEX = "F6F8FA"
BODY_FONT = "Calibri"

# verdict band colours
VERDICT_HEX = {
    "invest": "1C5E38", "back": "1C5E38", "proceed": "1C5E38",
    "conditional": "8A6D1F", "conditional invest": "8A6D1F", "hold": "8A6D1F",
    "pass": "8C2F2A", "decline": "8C2F2A",
}
# scorecard 0-5 colour ramp
SCORE_HEX = {0: "E8B4B0", 1: "E8B4B0", 2: "F0D9A8", 3: "F0D9A8",
             4: "C3DEC8", 5: "9ECBA8"}
SEVERITY_HEX = {"high": "8C2F2A", "medium": "8A6D1F", "med": "8A6D1F", "low": "6B6B6B"}

CHART_COLORS = ["#0F3D5C", "#4E8FB5", "#9EC5DC", "#C9A227", "#8C8C8C", "#B5563F"]


# ---------------------------------------------------------------- primitives
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def no_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def hairlines(table, color="C7CFD6"):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def cell_text(cell, text, size=8.5, bold=False, color=INK, align=None, space_after=1):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.text = ""
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(str(text))
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = BODY_FONT
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    r.font.name = BODY_FONT
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), ACCENT_HEX)
    bdr.append(bottom)
    pPr.append(bdr)
    return p


def body(doc, text, size=8.5, italic=False, color=INK, space_after=4, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.italic = italic
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = BODY_FONT
    return p


def bullets(doc, items, size=8.5, marker="\u25aa"):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.10)
        if isinstance(it, dict):
            r = p.add_run(f"{marker} " + it.get("point", ""))
            r.font.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = INK
            r.font.name = BODY_FONT
            if it.get("evidence"):
                r2 = p.add_run("  " + it["evidence"])
                r2.font.size = Pt(size)
                r2.font.color.rgb = MUTED
                r2.font.name = BODY_FONT
        else:
            r = p.add_run(f"{marker} " + str(it))
            r.font.size = Pt(size)
            r.font.color.rgb = INK
            r.font.name = BODY_FONT


def set_widths(table, inches):
    """python-docx autofit spreads columns evenly, which wastes the page on
    narrow numeric columns. Widths must be set on every cell, not just the col."""
    table.autofit = False
    layout = OxmlElement("w:tblLayout")      # without this, renderers re-fit anyway
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for i, w in enumerate(inches):          # tblGrid — what fixed layout reads
        if i < len(table.columns):
            table.columns[i].width = Inches(w)
    for row in table.rows:                  # and every cell, which Word prefers
        for i, w in enumerate(inches):
            if i < len(row.cells):
                row.cells[i].width = Inches(w)


def data_table(doc, spec, font=8):
    cols = spec.get("columns", [])
    rows = spec.get("rows", [])
    if not cols and not rows:
        return None
    ncols = len(cols) if cols else max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncols)
    t.autofit = True
    hairlines(t)
    if cols:
        hdr = t.add_row().cells
        for i, c in enumerate(cols):
            shade(hdr[i], HEADER_HEX)
            cell_text(hdr[i], c, size=font, bold=True,
                      align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i in range(ncols):
            val = row[i] if i < len(row) else ""
            if ri % 2 == 1:
                shade(cells[i], ZEBRA_HEX)
            cell_text(cells[i], val, size=font, bold=(i == 0),
                      align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    if spec.get("caption"):
        body(doc, spec["caption"], size=7, italic=True, color=MUTED, space_after=4)
    return t


def callout(doc, label, lines, fill=CALLOUT_HEX, label_color=ACCENT):
    t = doc.add_table(rows=1, cols=1)
    no_borders(t)
    c = t.rows[0].cells[0]
    shade(c, fill)
    p = c.paragraphs[0]
    p.text = ""
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run(label.upper())
    r.font.size = Pt(8.5)
    r.font.bold = True
    r.font.color.rgb = label_color
    r.font.name = BODY_FONT
    for line in lines:
        lp = c.add_paragraph()
        lp.paragraph_format.space_after = Pt(2)
        lp.paragraph_format.left_indent = Inches(0.08)
        rr = lp.add_run(line)
        rr.font.size = Pt(8.5)
        rr.font.color.rgb = INK
        rr.font.name = BODY_FONT
    c.paragraphs[-1].paragraph_format.space_after = Pt(5)
    body(doc, "", size=2, space_after=2)
    return t


# ---------------------------------------------------------------- charts
def make_chart(spec, outdir, idx):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    ctype = spec.get("type", "bar")
    x = spec.get("x", [])
    series = spec.get("series", [])

    # cohort retention triangle: rows = cohort labels, matrix = % retained per period
    if ctype == "heatmap":
        matrix = spec.get("matrix", [])
        ylabels = spec.get("y", [])
        fig, ax = plt.subplots(figsize=(6.4, 0.32 * max(len(matrix), 3) + 1.1), dpi=200)
        maxlen = max((len(r) for r in matrix), default=0)
        grid = [[(v if v is not None else float("nan")) for v in r] + [float("nan")] * (maxlen - len(r))
                for r in matrix]
        im = ax.imshow(grid, cmap="YlGnBu", aspect="auto", vmin=0,
                       vmax=spec.get("vmax", 100))
        for i, row in enumerate(grid):
            for j, v in enumerate(row):
                if v == v:  # not NaN
                    ax.text(j, i, f"{v:.0f}", ha="center", va="center", fontsize=6.5,
                            color="#1A1A1A" if v < spec.get("vmax", 100) * 0.6 else "white")
        ax.set_xticks(range(maxlen))
        ax.set_xticklabels(x if x else [f"M{k}" for k in range(maxlen)], fontsize=7)
        ax.set_yticks(range(len(ylabels)))
        ax.set_yticklabels(ylabels, fontsize=7)
        ax.set_title(spec.get("title", ""), fontsize=9, color="#0F3D5C",
                     fontweight="bold", loc="left", pad=8)
        fig.colorbar(im, ax=ax, shrink=0.7).ax.tick_params(labelsize=6)
        fig.tight_layout()
        path = os.path.join(outdir, f"chart_{idx}.png")
        fig.savefig(path, bbox_inches="tight")
        plt.close(fig)
        return path

    # market pyramid: tiers from apex (few actors, high value) to base
    if ctype == "pyramid":
        tiers = spec.get("tiers", [])        # [{label, actors, value, note}]
        n = len(tiers)
        fig, ax = plt.subplots(figsize=(6.4, 0.55 * n + 1.0), dpi=200)
        ax.set_xlim(0, 100)
        ax.set_ylim(0, n)
        ax.axis("off")
        for i, t in enumerate(tiers):
            y = n - 1 - i
            top_w = 34 + 66 * (i / max(n, 1))
            bot_w = 34 + 66 * ((i + 1) / max(n, 1))
            poly = [(50 - top_w / 2, y + 1), (50 + top_w / 2, y + 1),
                    (50 + bot_w / 2, y), (50 - bot_w / 2, y)]
            ax.add_patch(plt.Polygon(poly, closed=True,
                                     color=CHART_COLORS[i % len(CHART_COLORS)], alpha=0.9))
            ax.text(50, y + 0.5, t.get("label", ""), ha="center", va="center",
                    fontsize=7.5, color="white", fontweight="bold")
            ax.text(101, y + 0.68, t.get("actors", ""), ha="left", va="center",
                    fontsize=7, color="#1A1A1A")
            ax.text(101, y + 0.32, t.get("value", ""), ha="left", va="center",
                    fontsize=7, color="#6B6B6B")
        ax.text(101, n + 0.08, "Actors  /  Value", ha="left", va="bottom",
                fontsize=7, color="#0F3D5C", fontweight="bold")
        ax.set_ylim(0, n + 0.45)
        ax.set_title(spec.get("title", ""), fontsize=9, color="#0F3D5C",
                     fontweight="bold", loc="left", pad=4)
        ax.set_xlim(0, 150)
        fig.tight_layout()
        path = os.path.join(outdir, f"chart_{idx}.png")
        fig.savefig(path, bbox_inches="tight")
        plt.close(fig)
        return path

    # competitive positioning map: labelled points on two axes, size = scale
    if ctype == "positioning":
        pts = spec.get("points", [])         # [{name, x, y, size, highlight}]
        fig, ax = plt.subplots(figsize=(6.4, 4.2), dpi=200)
        for p in pts:
            hl = p.get("highlight", False)
            ax.scatter(p["x"], p["y"], s=max(40, float(p.get("size", 1)) * 60),
                       color="#B5563F" if hl else "#4E8FB5", alpha=0.85,
                       edgecolor="white", linewidth=0.8, zorder=3)
            ax.annotate(p.get("name", ""), (p["x"], p["y"]), textcoords="offset points",
                        xytext=(0, 8), ha="center", fontsize=6.8,
                        fontweight="bold" if hl else "normal",
                        color="#1A1A1A")
        ax.set_xlim(*spec.get("xlim", (0, 10)))
        ax.set_ylim(*spec.get("ylim", (0, 10)))
        ax.axhline((spec.get("ylim", (0, 10))[0] + spec.get("ylim", (0, 10))[1]) / 2,
                   color="#C7CFD6", linewidth=0.8, linestyle="--")
        ax.axvline((spec.get("xlim", (0, 10))[0] + spec.get("xlim", (0, 10))[1]) / 2,
                   color="#C7CFD6", linewidth=0.8, linestyle="--")
        ax.set_xlabel(spec.get("xlabel", ""), fontsize=7.5, color="#6B6B6B")
        ax.set_ylabel(spec.get("ylabel", ""), fontsize=7.5, color="#6B6B6B")
        ax.set_xticks([]); ax.set_yticks([])
        for q, (qx, qy) in zip(spec.get("quadrants", []),
                               [(0.02, 0.97), (0.98, 0.97), (0.02, 0.03), (0.98, 0.03)]):
            ax.text(qx, qy, q, transform=ax.transAxes, fontsize=6.5, color="#9EC5DC",
                    ha="left" if qx < 0.5 else "right", va="top" if qy > 0.5 else "bottom")
        ax.set_title(spec.get("title", ""), fontsize=9, color="#0F3D5C",
                     fontweight="bold", loc="left", pad=8)
        for sp in ("top", "right"):
            ax.spines[sp].set_visible(False)
        fig.tight_layout()
        path = os.path.join(outdir, f"chart_{idx}.png")
        fig.savefig(path, bbox_inches="tight")
        plt.close(fig)
        return path

    fig, ax = plt.subplots(figsize=(6.4, 2.6), dpi=200)
    if ctype == "line":
        for i, s in enumerate(series):
            ax.plot(x, s.get("values", []), marker="o", linewidth=2,
                    color=CHART_COLORS[i % len(CHART_COLORS)], label=s.get("name", ""))
        if spec.get("floor") is not None:      # e.g. the retention plateau you claim
            ax.axhline(spec["floor"], color="#B5563F", linestyle="--", linewidth=1)
    elif ctype == "barh":
        s = series[0] if series else {"values": []}
        ax.barh(x, s.get("values", []), color=CHART_COLORS[0], height=0.6)
        ax.invert_yaxis()
    elif ctype == "waterfall":
        # NRR bridge: starting ARR, +expansion, -churn, -contraction, ending ARR
        vals = series[0].get("values", []) if series else []
        run, bottoms, heights, colors = 0, [], [], []
        for i, v in enumerate(vals):
            terminal = i in (0, len(vals) - 1)
            bottoms.append(0 if terminal else (run if v >= 0 else run + v))
            heights.append(v if terminal else abs(v))
            colors.append(CHART_COLORS[0] if terminal else
                          ("#4E8FB5" if v >= 0 else "#B5563F"))
            run = v if terminal else run + v
        ax.bar(range(len(vals)), heights, bottom=bottoms, color=colors, width=0.6)
        ax.set_xticks(range(len(vals)))
        ax.set_xticklabels(x, fontsize=7)
    else:  # bar / grouped_bar
        n = max(len(series), 1)
        width = 0.8 / n
        pos = range(len(x))
        for i, s in enumerate(series):
            offs = [p - 0.4 + width / 2 + i * width for p in pos]
            ax.bar(offs, s.get("values", []), width=width,
                   color=CHART_COLORS[i % len(CHART_COLORS)], label=s.get("name", ""))
        ax.set_xticks(list(pos))
        ax.set_xticklabels(x)

    ax.set_title(spec.get("title", ""), fontsize=9, color="#0F3D5C",
                 fontweight="bold", loc="left", pad=8)
    if spec.get("ylabel"):
        ax.set_ylabel(spec["ylabel"], fontsize=7.5, color="#6B6B6B")
    ax.tick_params(labelsize=7.5, colors="#4A4A4A")
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    for sp in ("left", "bottom"):
        ax.spines[sp].set_color("#C7CFD6")
    ax.grid(axis="y", color="#E6EAEE", linewidth=0.7)
    ax.set_axisbelow(True)
    if len([s for s in series if s.get("name")]) > 1:
        ax.legend(fontsize=7, frameon=False)
    fig.tight_layout()
    path = os.path.join(outdir, f"chart_{idx}.png")
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def add_chart(doc, spec, outdir, counter):
    try:
        path = make_chart(spec, outdir, counter)
    except Exception as e:            # a broken chart must never kill the document
        body(doc, f"[chart omitted: {e}]", size=7, italic=True, color=MUTED)
        return
    doc.add_picture(path, width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if spec.get("source"):
        body(doc, spec["source"], size=7, italic=True, color=MUTED, space_after=4)


# ---------------------------------------------------------------- sections
def build_header(doc, m):
    h = m.get("header", {})
    t = doc.add_table(rows=1, cols=1)
    no_borders(t)
    c = t.rows[0].cells[0]
    shade(c, ACCENT_HEX)
    c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = c.paragraphs[0]
    p.text = ""
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(m.get("company", "Company"))
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = BODY_FONT
    sub = p.add_run("   " + m.get("classification", ""))
    sub.font.size = Pt(9)
    sub.font.color.rgb = RGBColor(0xB8, 0xD0, 0xE0)
    sub.font.name = BODY_FONT

    line = c.add_paragraph()
    line.paragraph_format.space_after = Pt(5)
    bits = [h.get("product"), h.get("stage"), h.get("round"), h.get("ask"),
            h.get("arr"), h.get("customers"), h.get("hq")]
    r2 = line.add_run("  \u00b7  ".join([b for b in bits if b]))
    r2.font.size = Pt(8)
    r2.font.color.rgb = RGBColor(0xD4, 0xE2, 0xEB)
    r2.font.name = BODY_FONT

    label = m.get("doc_label", "Product due diligence")
    if h.get("date"):
        d = body(doc, f"{label} \u00b7 {h['date']}", size=7, italic=True,
                 color=MUTED, space_after=6)
        d.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def build_verdict(doc, m):
    """Full-width coloured band: the call, in one line, before any analysis."""
    v = m.get("verdict")
    if not v:
        return
    call = v.get("call", "")
    fill = VERDICT_HEX.get(call.strip().lower(), ACCENT_HEX)
    t = doc.add_table(rows=1, cols=1)
    no_borders(t)
    c = t.rows[0].cells[0]
    shade(c, fill)
    p = c.paragraphs[0]
    p.text = ""
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(call.upper())
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = BODY_FONT
    if v.get("conviction"):
        cv = p.add_run(f"   conviction: {v['conviction']}")
        cv.font.size = Pt(8)
        cv.font.color.rgb = RGBColor(0xE0, 0xE8, 0xEE)
        cv.font.name = BODY_FONT
    if v.get("one_liner"):
        lp = c.add_paragraph()
        lp.paragraph_format.space_after = Pt(5)
        rr = lp.add_run(v["one_liner"])
        rr.font.size = Pt(9)
        rr.font.color.rgb = WHITE
        rr.font.name = BODY_FONT
    body(doc, "", size=2, space_after=2)


def build_read(doc, m):
    items = m.get("the_read", [])
    if items:
        callout(doc, "The read", ["\u25aa " + i for i in items])


def build_scorecard(doc, m):
    """Weighted 0-5 grid. The colour is the signal; the finding is the argument."""
    sc = m.get("scorecard")
    if not sc:
        return
    rows = sc.get("rows", [])
    heading(doc, sc.get("title", "Product scorecard"))
    t = doc.add_table(rows=0, cols=5)
    t.autofit = True
    hairlines(t)
    hdr = t.add_row().cells
    for i, c in enumerate(["Module", "Score", "Wt", "Finding", "Conf."]):
        shade(hdr[i], HEADER_HEX)
        cell_text(hdr[i], c, size=8, bold=True,
                  align=WD_ALIGN_PARAGRAPH.LEFT if i in (0, 3) else WD_ALIGN_PARAGRAPH.CENTER)
    total_w, total_s = 0.0, 0.0
    for r in rows:
        score = r.get("score")
        w = float(r.get("weight", 1) or 0)
        if isinstance(score, (int, float)):
            total_w += w
            total_s += w * float(score)
        cells = t.add_row().cells
        cell_text(cells[0], r.get("module", ""), size=8, bold=True)
        cell_text(cells[1], f"{score}/5" if score is not None else "n/a", size=8,
                  bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        if isinstance(score, (int, float)):
            shade(cells[1], SCORE_HEX.get(int(round(score)), ZEBRA_HEX))
        cell_text(cells[2], f"{w:g}", size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(cells[3], r.get("finding", ""), size=8)
        cell_text(cells[4], r.get("confidence", ""), size=8,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    if total_w:
        cells = t.add_row().cells
        for i in range(5):
            shade(cells[i], HEADER_HEX)
        cell_text(cells[0], "Weighted total", size=8, bold=True)
        cell_text(cells[1], f"{total_s / total_w:.1f}/5", size=8, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(cells[2], f"{total_w:g}", size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(cells[3], sc.get("total_note", ""), size=8)
        cell_text(cells[4], "", size=8)
    set_widths(t, [1.45, 0.55, 0.35, 3.45, 0.5])
    if sc.get("caption"):
        body(doc, sc["caption"], size=7, italic=True, color=MUTED)


def build_sections(doc, m):
    """Generic analytical body: each entry may carry prose, bullets, a table, a chart."""
    for s in m.get("sections", []):
        heading(doc, s.get("heading", ""))
        if s.get("prose"):
            body(doc, s["prose"])
        if s.get("bullets"):
            bullets(doc, s["bullets"])
        if s.get("table"):
            data_table(doc, s["table"])
        for extra in s.get("tables", []):
            data_table(doc, extra)
        if s.get("so_what"):
            body(doc, "So what: " + s["so_what"], italic=True, color=ACCENT)


def build_red_flags(doc, m):
    flags = m.get("red_flags", [])
    if not flags:
        return
    heading(doc, "Red flags")
    t = doc.add_table(rows=0, cols=4)
    t.autofit = True
    hairlines(t)
    hdr = t.add_row().cells
    for i, c in enumerate(["Flag", "Sev.", "Evidence", "Mitigation / what would change our mind"]):
        shade(hdr[i], HEADER_HEX)
        cell_text(hdr[i], c, size=8, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER if i == 1 else WD_ALIGN_PARAGRAPH.LEFT)
    for ri, f in enumerate(flags):
        cells = t.add_row().cells
        if ri % 2 == 1:
            for c in cells:
                shade(c, ZEBRA_HEX)
        cell_text(cells[0], f.get("flag", ""), size=8, bold=True)
        sev = str(f.get("severity", "")).lower()
        cell_text(cells[1], sev.upper()[:4], size=7.5, bold=True,
                  color=RGBColor.from_string(SEVERITY_HEX.get(sev, "6B6B6B")),
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(cells[2], f.get("evidence", ""), size=8)
        cell_text(cells[3], f.get("mitigation", ""), size=8)
    set_widths(t, [1.5, 0.45, 2.2, 2.15])


def build_questions(doc, m):
    qs = m.get("open_questions", [])
    if not qs:
        return
    heading(doc, m.get("questions_heading", "Questions for management"))
    for q in qs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.1)
        r = p.add_run("? " + q)
        r.font.size = Pt(8.5)
        r.font.color.rgb = INK
        r.font.name = BODY_FONT


def build_data_gaps(doc, m):
    """Appendix: what raw data is still needed, from whom, and how to get it."""
    gaps = m.get("data_gaps", [])
    if not gaps:
        return
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("APPENDIX \u2014 RAW DATA STILL REQUIRED")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    r.font.name = BODY_FONT
    body(doc, m.get("data_gaps_intro",
                    "Findings above are tagged by confidence. The items below would move "
                    "[B]/[C] claims to [A] or close an open module. Priority 1 = needed before "
                    "any recommendation; 2 = needed before close; 3 = post-close."),
         size=8, italic=True, color=MUTED)
    t = doc.add_table(rows=0, cols=6)
    hairlines(t)
    hdr = t.add_row().cells
    for i, c in enumerate(["Pri.", "What we need", "Why / which module", "From whom",
                           "How to obtain", "Status"]):
        shade(hdr[i], HEADER_HEX)
        cell_text(hdr[i], c, size=7.5, bold=True)
    for ri, g in enumerate(gaps):
        cells = t.add_row().cells
        if ri % 2 == 1:
            for c in cells:
                shade(c, ZEBRA_HEX)
        cell_text(cells[0], str(g.get("priority", "")), size=7.5, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(cells[1], g.get("what", ""), size=7.5, bold=True)
        cell_text(cells[2], g.get("why", ""), size=7.5)
        cell_text(cells[3], g.get("from_whom", ""), size=7.5)
        cell_text(cells[4], g.get("how", ""), size=7.5)
        cell_text(cells[5], g.get("status", "open"), size=7.5)
    set_widths(t, [0.35, 1.45, 1.5, 1.15, 1.6, 0.55])


# ---------------------------------------------------------------- assembly
def build(doc_json, out_path):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.5)
    sec.bottom_margin = Inches(0.5)
    sec.left_margin = Inches(0.6)
    sec.right_margin = Inches(0.6)

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(8.5)
    style.paragraph_format.space_after = Pt(3)

    tmpdir = tempfile.mkdtemp(prefix="pdd_charts_")
    counter = [0]

    build_header(doc, doc_json)
    build_verdict(doc, doc_json)
    build_read(doc, doc_json)
    build_scorecard(doc, doc_json)
    build_sections(doc, doc_json)
    build_red_flags(doc, doc_json)
    build_questions(doc, doc_json)

    exhibits = doc_json.get("exhibits", [])
    if exhibits:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run("EXHIBITS")
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = ACCENT
        r.font.name = BODY_FONT
        for i, ex in enumerate(exhibits):
            if i > 0:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            heading(doc, f"Exhibit {ex.get('id', chr(65 + i))} \u2014 {ex.get('title', '')}")
            if ex.get("prose"):
                body(doc, ex["prose"])
            if ex.get("chart"):
                counter[0] += 1
                add_chart(doc, ex["chart"], tmpdir, counter[0])
            for ch in ex.get("charts", []):
                counter[0] += 1
                add_chart(doc, ch, tmpdir, counter[0])
            if ex.get("bullets"):
                bullets(doc, ex["bullets"], size=8)
            if ex.get("table"):
                data_table(doc, ex["table"], font=7.5)
            for extra in ex.get("tables", []):
                data_table(doc, extra, font=7.5)
            if ex.get("notes"):
                body(doc, ex["notes"], size=7.5, italic=True, color=MUTED)

    build_data_gaps(doc, doc_json)

    body(doc, doc_json.get(
        "confidence_note",
        "Confidence tags: [A] verified in data room / audited  \u00b7  [B] management-stated or "
        "credible third party  \u00b7  [C] estimated, inferred or single-source."),
        size=7, italic=True, color=MUTED, space_after=0)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("doc_json")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    with open(args.doc_json) as f:
        data = json.load(f)
    print(f"Wrote {build(data, args.out)}")


if __name__ == "__main__":
    main()
