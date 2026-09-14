#!/usr/bin/env python3
"""Render outreach.json into a Word outreach plan: at-a-glance table + one section per target.

Usage:
    python3 build_outreach_doc.py outreach.json --out /path/to/NAME_outreach.docx

See references/outreach-schema.md.
"""

import argparse
import json
import os

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ---------------------------------------------------------------- palette
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
ACCENT = RGBColor(0x0F, 0x3D, 0x5C)
RUST = RGBColor(0xB5, 0x56, 0x3F)
GREEN = RGBColor(0x2E, 0x6B, 0x3E)
AMBER = RGBColor(0x8A, 0x6D, 0x1F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_HEX = "0F3D5C"
CALLOUT_HEX = "EEF3F7"
HEADER_HEX = "DCE5EC"
ZEBRA_HEX = "F6F8FA"
MSG_HEX = "FBF3E9"        # message bodies — the payload, warm and obvious
EVENT_HEX = "EDF2ED"      # where they'll physically be
TIP_HEX = "F3F1EC"
BODY_FONT = "Calibri"


# ---------------------------------------------------------------- helpers
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _borders(table, val, sz="4", color="C7CFD6"):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), sz)
        if val != "none":
            el.set(qn("w:color"), color)
        b.append(el)
    table._tbl.tblPr.append(b)


def no_borders(t):
    _borders(t, "none", "0")


def hairlines(t, color="C7CFD6"):
    _borders(t, "single", "4", color)


def run(p, text, size=8.5, bold=False, italic=False, color=INK):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = BODY_FONT
    return r


def para(c, space_after=3, space_before=0, indent=None, first=False):
    if first and c.paragraphs:
        p = c.paragraphs[0]
        p.text = ""
    else:
        p = c.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if indent is not None:
        pf.left_indent = Inches(indent)
    return p


def page_break_before(doc):
    """New page with no blank page left behind (see build_people_doc.py for why)."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.page_break_before = True
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    p.add_run("").font.size = Pt(1)
    return p


def section_label(c, text, color=ACCENT, size=7.5, space_before=4):
    p = para(c, space_after=1, space_before=space_before)
    run(p, text.upper(), size=size, bold=True, color=color)
    return p


def odds_color(odds):
    o = (odds or "").strip().lower()
    if o in ("best", "high", "strong"):
        return GREEN
    if o in ("medium", "moderate", "possible"):
        return AMBER
    return RUST


def data_table(doc, spec, font=8):
    cols = spec.get("columns", [])
    rows = spec.get("rows", [])
    if not rows:
        return
    t = doc.add_table(rows=0, cols=len(cols))
    hairlines(t)
    hdr = t.add_row().cells
    for i, c in enumerate(cols):
        shade(hdr[i], HEADER_HEX)
        run(para(hdr[i], space_after=1, first=True), c, size=font, bold=True)
    for ri, r in enumerate(rows):
        cells = t.add_row().cells
        for i in range(len(cols)):
            if ri % 2 == 1:
                shade(cells[i], ZEBRA_HEX)
            run(para(cells[i], space_after=1, first=True),
                r[i] if i < len(r) else "", size=font, bold=(i == 0))


# ---------------------------------------------------------------- sections
def build_header(doc, m):
    t = doc.add_table(rows=1, cols=1)
    no_borders(t)
    c = t.rows[0].cells[0]
    shade(c, ACCENT_HEX)
    c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = para(c, space_before=4, space_after=1, first=True)
    run(p, m.get("title", "Outreach plan"), size=16, bold=True, color=WHITE)
    run(p, "   outreach plan", size=9, color=RGBColor(0xB8, 0xD0, 0xE0))
    if m.get("angle"):
        p2 = para(c, space_after=5)
        run(p2, "THE ANGLE  ", size=7.5, bold=True, color=RGBColor(0x9E, 0xC5, 0xDC))
        run(p2, m["angle"], size=8.5, color=RGBColor(0xEC, 0xF2, 0xF6))
    d = para(doc, space_after=6)
    d.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(d, m.get("date", ""), size=7, italic=True, color=MUTED)


def build_routes(doc, routes):
    if not routes:
        return
    section_label(doc, "Routes in — ranked by what actually lands")
    t = doc.add_table(rows=0, cols=3)
    hairlines(t)
    hdr = t.add_row().cells
    for i, h in enumerate(["Route", "Odds", "Why"]):
        shade(hdr[i], HEADER_HEX)
        run(para(hdr[i], space_after=1, first=True), h, size=7.5, bold=True)
    for ri, r in enumerate(routes):
        cells = t.add_row().cells
        for i in range(3):
            if ri % 2 == 1:
                shade(cells[i], ZEBRA_HEX)
        run(para(cells[0], space_after=1, first=True), r.get("route", ""), size=8, bold=True)
        p = para(cells[1], space_after=1, first=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, r.get("odds", ""), size=8, bold=True, color=odds_color(r.get("odds")))
        run(para(cells[2], space_after=1, first=True), r.get("why", ""), size=7.5, color=MUTED)


def build_events(doc, events):
    if not events:
        return
    section_label(doc, "Where they will physically be", color=GREEN)
    for e in events:
        t = doc.add_table(rows=1, cols=1)
        hairlines(t, color="BFD2BF")
        c = t.rows[0].cells[0]
        shade(c, EVENT_HEX)
        p = para(c, space_before=2, space_after=1, first=True)
        run(p, e.get("event", ""), size=9, bold=True, color=ACCENT)
        bits = [e.get("date"), e.get("city")]
        run(p, "   " + "  ·  ".join([b for b in bits if b]), size=7.5, color=MUTED)
        for label, key in [("Their role", "their_role"), ("Access", "access"),
                           ("The approach", "approach")]:
            if e.get(key):
                pp = para(c, space_after=1, indent=0.04)
                run(pp, label + ": ", size=7.5, bold=True, color=MUTED)
                run(pp, e[key], size=8)
        if e.get("source"):
            sp = para(c, space_after=3, indent=0.04)
            run(sp, e["source"], size=7, italic=True, color=MUTED)
        para(doc, space_after=2)


def build_sequence(doc, seq):
    if not seq:
        return
    section_label(doc, "The sequence")
    for step in seq:
        # step header strip
        t = doc.add_table(rows=1, cols=1)
        no_borders(t)
        c = t.rows[0].cells[0]
        shade(c, CALLOUT_HEX)
        p = para(c, space_before=2, space_after=2, first=True)
        run(p, (step.get("day") or "") + "   ", size=9, bold=True, color=RUST)
        run(p, step.get("purpose", ""), size=8.5, bold=True)
        run(p, "    via " + step.get("channel", ""), size=7.5, color=MUTED)

        if step.get("destination"):
            dp = para(doc, space_after=1, indent=0.04)
            run(dp, "→ ", size=8, bold=True, color=ACCENT)
            run(dp, step["destination"], size=8, bold=True, color=ACCENT)

        # the message itself
        mt = doc.add_table(rows=1, cols=1)
        hairlines(mt, color="D9C4A8")
        mc = mt.rows[0].cells[0]
        shade(mc, MSG_HEX)
        if step.get("subject"):
            sp = para(mc, space_before=2, space_after=2, first=True)
            run(sp, "Subject: ", size=7.5, bold=True, color=MUTED)
            run(sp, step["subject"], size=8.5, bold=True)
            bp = para(mc, space_after=2)
        else:
            bp = para(mc, space_before=2, space_after=2, first=True)
        run(bp, step.get("body", ""), size=8.5)
        if step.get("note"):
            np_ = para(mc, space_after=3)
            run(np_, "Note: ", size=7.5, bold=True, color=MUTED)
            run(np_, step["note"], size=7.5, italic=True, color=MUTED)
        para(doc, space_after=2)


def build_tip(doc, tip):
    if not tip:
        return
    t = doc.add_table(rows=1, cols=1)
    hairlines(t, color="D2CBBC")
    c = t.rows[0].cells[0]
    shade(c, TIP_HEX)
    p = para(c, space_before=2, space_after=1, first=True)
    run(p, "TIP THAT RAISES CONVERSION   ", size=7.5, bold=True, color=RUST)
    run(p, tip.get("tip", ""), size=8.5, bold=True)
    if tip.get("why"):
        wp = para(c, space_after=3)
        run(wp, "Why it works: ", size=7.5, bold=True, color=MUTED)
        run(wp, tip["why"], size=8, italic=True, color=MUTED)


def build_target(doc, tgt, first=False):
    if not first:
        page_break_before(doc)
    # name strip
    t = doc.add_table(rows=1, cols=1)
    no_borders(t)
    c = t.rows[0].cells[0]
    shade(c, ACCENT_HEX)
    p = para(c, space_before=3, space_after=2, first=True)
    run(p, tgt.get("name", ""), size=13, bold=True, color=WHITE)
    run(p, "   " + tgt.get("title", ""), size=8, color=RGBColor(0xB8, 0xD0, 0xE0))

    if tgt.get("objective"):
        op = para(doc, space_after=2, space_before=3)
        run(op, "OBJECTIVE   ", size=7.5, bold=True, color=RUST)
        run(op, tgt["objective"], size=9, bold=True)
    if tgt.get("read"):
        rp = para(doc, space_after=3)
        run(rp, tgt["read"], size=8.5, color=INK)

    build_routes(doc, tgt.get("routes"))
    build_events(doc, tgt.get("where_they_will_be"))
    build_sequence(doc, tgt.get("sequence"))
    build_tip(doc, tgt.get("tip"))


# ---------------------------------------------------------------- main
def build(m, out_path):
    doc = Document()
    s = doc.sections[0]
    s.top_margin = Inches(0.5)
    s.bottom_margin = Inches(0.5)
    s.left_margin = Inches(0.6)
    s.right_margin = Inches(0.6)
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(8.5)
    st.paragraph_format.space_after = Pt(3)

    build_header(doc, m)

    if m.get("at_a_glance"):
        section_label(doc, "The plan at a glance", size=9)
        data_table(doc, m["at_a_glance"])

    for i, tgt in enumerate(m.get("targets", [])):
        build_target(doc, tgt, first=False)

    if m.get("source_log"):
        page_break_before(doc)
        section_label(doc, "Source log", size=9)
        data_table(doc, m["source_log"], font=7.5)

    p = para(doc, space_after=0)
    run(p, m.get("footer_note",
                 "Email addresses are inferred from the company's observed pattern and are not "
                 "verified — the % is the confidence in the pattern. [Brackets] mark what only you "
                 "can supply: never send them unfilled. Public professional settings only."),
        size=7, italic=True, color=MUTED)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("outreach_json")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    with open(a.outreach_json) as f:
        m = json.load(f)
    print("Wrote " + build(m, a.out))


if __name__ == "__main__":
    main()
